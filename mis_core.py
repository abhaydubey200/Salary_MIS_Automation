"""
APIS India Limited — Salary MIS core logic.

Faithful Python port of the parsing / modelling / cross-check logic in
APIS_India_Limited___Salary_MIS_Automation.html.

Design notes
------------
* All money handling mirrors the JS `n()` helper: strings with thousands
  separators are cleaned, blanks and non-numeric values become 0.
* Column discovery mirrors the JS `FIND` regex table and `mapCols()`:
  the sheet layout is located by HEADER NAME, not fixed position, so the
  same monthly template keeps working if a column shifts.
* The MIS model (`build_model`) reproduces `model()` exactly: the HO
  department order, Directors block, Sales channels, the manual memo
  block, and the two formula sets ("tpl" = senior's template as given,
  "fix" = corrected / Variable counted once).
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from typing import Any

import openpyxl

# ---------------------------------------------------------------------------
# Constants (mirror the JS)
# ---------------------------------------------------------------------------

HO_ORDER = [
    "Admin & Facilities", "Business Excellence", "Company Secretarial", "Export",
    "Finance & Accounts", "Information Technology", "Manufacturing", "Marketing",
    "New Product Development", "People & Culture", "Procurement", "Sales",
    "Supply Chain Management",
]

MON = ["January", "February", "March", "April", "May", "June", "July",
       "August", "September", "October", "November", "December"]

HOLD_RE = re.compile(r"resign|abscond|hold|\batl\b|terminat|accident", re.I)

# Column-finding regexes, keyed exactly like the JS FIND object.
FIND = {
    "var":   re.compile(r"^var[a-z]*le", re.I),
    "arr":   re.compile(r"arrear", re.I),
    "gross": re.compile(r"^(total gross earned|fixed gross earning)", re.I),
    "reimb": re.compile(r"gross\s*reimb", re.I),
    "tds":   re.compile(r"^tds$", re.I),
    "pf":    re.compile(r"^pf\s*employee", re.I),
    "esic":  re.compile(r"^esic\s*employee", re.I),
    "loan":  re.compile(r"^loan amount", re.I),
    "adv":   re.compile(r"^sal(ary)?\s*advance", re.I),
    "oth":   re.compile(r"^other deduction", re.I),
    "nps":   re.compile(r"^nps$", re.I),
    "ded":   re.compile(r"^total deductions", re.I),
    "net":   re.compile(r"^salary net pay", re.I),
    "epf":   re.compile(r"^employer pf", re.I),
    "eesic": re.compile(r"^(employer|emloyer) esic", re.I),
    "rem":   re.compile(r"^remarks", re.I),
}

# Numeric keys carried per employee (used for junk-text detection in checks).
_EMP_NUM_KEYS = ["var", "arr", "gross", "reimb", "tds", "pf", "esic",
                 "loan", "adv", "oth", "nps", "ded", "net", "epf", "eesic"]

# Human-readable names for the junk-text check.
NL = {"var": "Variable", "arr": "Arrear", "gross": "Gross", "reimb": "Reimbursement",
      "tds": "TDS", "pf": "PF", "esic": "Employee ESIC", "loan": "Loan",
      "adv": "Advance", "oth": "Other deduction", "nps": "NPS",
      "ded": "Total deductions", "net": "Net pay", "epf": "Employer PF",
      "eesic": "Employer ESIC"}


# ---------------------------------------------------------------------------
# Small helpers
# ---------------------------------------------------------------------------

def n(x: Any) -> float:
    """Port of the JS `n()` coercion: tolerant numeric parse, blanks/garbage -> 0."""
    if isinstance(x, bool):
        return 0.0
    if isinstance(x, (int, float)):
        return float(x) if _is_finite(x) else 0.0
    if x is None or x == "":
        return 0.0
    try:
        v = float(str(x).replace(",", ""))
        return v if _is_finite(v) else 0.0
    except ValueError:
        return 0.0


def _is_finite(v: float) -> bool:
    return v == v and v not in (float("inf"), float("-inf"))


def fmt_in(value: float) -> str:
    """Indian-grouping integer formatter, matching the app's toLocaleString('en-IN')."""
    num = round(n(value))
    neg = num < 0
    s = str(abs(int(num)))
    if len(s) <= 3:
        out = s
    else:
        head, tail = s[:-3], s[-3:]
        parts = []
        while len(head) > 2:
            parts.insert(0, head[-2:])
            head = head[:-2]
        if head:
            parts.insert(0, head)
        out = ",".join(parts) + "," + tail
    return ("-" if neg else "") + out


def key(s: str) -> str:
    return str(s).strip().lower()


def detect_month(name: str) -> str:
    """Port of JS detectMonth(): returns 'YYYY-MM' or ''."""
    m = re.search(r"(20\d\d)[._\- ]?(0[1-9]|1[0-2])", name)
    if m:
        return f"{m.group(1)}-{m.group(2)}"
    for i in range(12):
        r = re.search(MON[i][:3] + r"[a-z]*[\s,_\-']*(20)?(\d\d)", name, re.I)
        if r:
            return f"20{r.group(2)}-{str(i + 1).zfill(2)}"
    return ""


# ---------------------------------------------------------------------------
# Decryption + reading
# ---------------------------------------------------------------------------

class SheetError(Exception):
    """Raised with a user-facing message, mirroring the JS thrown errors."""


def _maybe_decrypt(data: bytes, passwords: list[str]) -> bytes:
    """If `data` is a CFB (encrypted) OOXML package, decrypt with the first
    working password. Mirrors the JS unlock() path (which handled modern
    AES-protected .xlsx). Returns plaintext .xlsx bytes."""
    if len(data) >= 2 and data[0] == 0xD0 and data[1] == 0xCF:
        import msoffcrypto  # imported lazily so the app loads without it
        pws = [p for p in passwords if p]
        if not pws:
            raise SheetError(
                'This file is password-protected. Enter its password in the '
                '"File password" box and submit again (or remove the password in Excel).')
        last = None
        for p in pws:
            try:
                off = msoffcrypto.OfficeFile(io.BytesIO(data))
                off.load_key(password=p)
                out = io.BytesIO()
                off.decrypt(out)
                return out.getvalue()
            except Exception as e:  # noqa: BLE001
                last = e
        raise SheetError(
            f"Could not unlock the file ({last}). Check the password, or "
            "remove it in Excel and upload again.")
    return data


def _pick_sheet(names: list[str], kind: str) -> str:
    if kind == "ho":
        for nm in names:
            if not re.search(r"retainer", nm, re.I):
                return nm
        return names[0]
    # sales: prefer a sheet named like "APIS", else first non-retainer
    for nm in names:
        if re.search(r"apis", nm, re.I):
            return nm
    for nm in names:
        if not re.search(r"retainer", nm, re.I):
            return nm
    return names[0]


@dataclass
class Employee:
    code: str
    name: str
    dept: str
    rem: str
    fixed: float          # sum of the 5 earned components before Variable (Basic..Special)
    flag: bool            # ESIC flag == YES
    hasFlag: bool         # both flag and ac columns present
    ac: float             # "gross earn" column in the 20..40 band (for ESIC rule check)
    ap: float             # first fixed component (Basic) — ESIC basis variant
    at: float             # last fixed component (Special allowance) — ESIC basis variant
    # numeric fields:
    var: float = 0.0
    arr: float = 0.0
    gross: float = 0.0
    reimb: float = 0.0
    tds: float = 0.0
    pf: float = 0.0
    esic: float = 0.0
    loan: float = 0.0
    adv: float = 0.0
    oth: float = 0.0
    nps: float = 0.0
    ded: float = 0.0
    net: float = 0.0
    epf: float = 0.0
    eesic: float = 0.0


@dataclass
class ParsedSheet:
    emps: list[Employee]
    warn: list[str]
    sheet: str
    fileName: str
    tot: dict | None
    junk: list[dict]


def _map_cols(h: list) -> tuple[dict, list[str]]:
    """Port of JS mapCols(). `h` is the header row (0-indexed list)."""
    def cell(j):
        return "" if j < 0 or j >= len(h) or h[j] is None else str(h[j]).strip()

    col = {"code": 1, "name": 2, "dept": 3}
    miss: list[str] = []
    for k, rx in FIND.items():
        found = -1
        for j in range(len(h)):
            if j >= 38 and rx.match(cell(j)):
                found = j
                break
        col[k] = found
        if found < 0 and k != "rem":
            miss.append(k)

    col["fx"] = [col["var"] - d for d in (5, 4, 3, 2, 1)]

    col["flag"] = -1
    for j in range(min(20, len(h))):
        if re.match(r"^esic$", cell(j), re.I):
            col["flag"] = j
            break
    col["ac"] = -1
    for j in range(20, min(40, len(h))):
        if re.search(r"gross\s*earn", cell(j), re.I):
            col["ac"] = j
            break

    if col["var"] >= 0:
        first_fx = cell(col["fx"][0])
        last_fx = cell(col["fx"][4])
        if not re.search(r"basic", first_fx, re.I) or not re.search(r"special", last_fx, re.I):
            miss.append("earned components (Basic..Special) before Variable")
    if (not re.search(r"employee code", cell(1), re.I)
            or not re.search(r"name", cell(2), re.I)
            or not re.search(r"dep", cell(3), re.I)):
        miss.append("Employee Code / Name / Department")
    return col, miss


def load_sheet(data: bytes, filename: str, kind: str, passwords: list[str]) -> ParsedSheet:
    """Port of JS load(). Returns a ParsedSheet or raises SheetError."""
    plain = _maybe_decrypt(data, passwords)

    wb = openpyxl.load_workbook(io.BytesIO(plain), data_only=True, read_only=True)
    sn = _pick_sheet(wb.sheetnames, kind)
    ws = wb[sn]

    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    if not rows:
        raise SheetError("The sheet is empty.")

    # header row: first of first 8 rows whose column B matches "Employee Code"
    h_idx = -1
    for i in range(min(8, len(rows))):
        b = rows[i][1] if len(rows[i]) > 1 else None
        if re.search(r"employee\s*code", str(b or ""), re.I):
            h_idx = i
            break
    if h_idx < 0:
        raise SheetError('Header "Employee Code" not found in column B (first 8 rows). '
                         "Is this the right file/sheet?")

    header = rows[h_idx]
    col, miss = _map_cols(header)
    if miss:
        raise SheetError("Required column(s) not found by header name: " + ", ".join(miss) +
                         ". The sheet layout has changed — please check the file.")

    warn: list[str] = []
    if col["rem"] < 0:
        warn.append('No "Remarks" column found — hold list cannot be pre-selected')

    def g(row, j):
        return row[j] if (0 <= j < len(row)) else None

    emps: list[Employee] = []
    junk: list[dict] = []
    last_i = h_idx
    for i in range(h_idx + 1, len(rows)):
        r = rows[i]
        code_raw = str(g(r, col["code"]) or "")
        if not re.match(r"^[A-Za-z]{2,4}\d+", code_raw):
            continue
        if not g(r, col["dept"]):
            continue
        fx = col["fx"]
        e = Employee(
            code=code_raw.strip(),
            name=str(g(r, col["name"]) or "").strip(),
            dept=str(g(r, col["dept"])).strip(),
            rem="" if col["rem"] < 0 else str(g(r, col["rem"]) if g(r, col["rem"]) is not None else "").strip(),
            fixed=sum(n(g(r, c)) for c in fx),
            flag=(col["flag"] >= 0 and re.match(r"^yes$", str(g(r, col["flag"]) or "").strip(), re.I) is not None),
            hasFlag=(col["flag"] >= 0 and col["ac"] >= 0),
            ac=(n(g(r, col["ac"])) if col["ac"] >= 0 else 0.0),
            ap=n(g(r, fx[0])),
            at=n(g(r, fx[4])),
        )
        for k in _EMP_NUM_KEYS:
            rw = g(r, col[k])
            setattr(e, k, n(rw))
            if isinstance(rw, str) and rw.strip() != "":
                try:
                    float(rw.replace(",", ""))
                except ValueError:
                    junk.append({"code": e.code, "k": k, "v": rw.strip()})
        emps.append(e)
        last_i = i

    if not emps:
        raise SheetError("No employee rows found.")

    prefix = "HO" if kind == "ho" else "SL"
    pre = sum(1 for e in emps if re.match("^" + prefix, e.code, re.I))
    if pre < len(emps) * 0.8:
        if kind == "ho":
            raise SheetError("This does not look like the HO sheet (employee codes should "
                             "start with HO). Did you upload the Sales file here?")
        raise SheetError("This does not look like the Sales sheet (employee codes should "
                         "start with SL). Did you upload the HO file here?")

    # source total row (first blank-code row after data with a positive gross)
    tot = None
    for i in range(last_i + 1, min(len(rows), last_i + 8)):
        r = rows[i]
        code_blank = not str(g(r, col["code"]) if g(r, col["code"]) is not None else "").strip()
        gv = g(r, col["gross"])
        if code_blank and isinstance(gv, (int, float)) and gv > 0:
            epf_v = g(r, col["epf"])
            eesic_v = g(r, col["eesic"])
            tot = {
                "gross": n(gv), "ded": n(g(r, col["ded"])), "net": n(g(r, col["net"])),
                "epf": epf_v if isinstance(epf_v, (int, float)) else None,
                "eesic": eesic_v if isinstance(eesic_v, (int, float)) else None,
                "row": i + 1,
            }
            break

    return ParsedSheet(emps=emps, warn=warn, sheet=sn, fileName=filename, tot=tot, junk=junk)


# ---------------------------------------------------------------------------
# Aggregation + model
# ---------------------------------------------------------------------------

# Aggregation keys used by agg()/sumO (mirror JS KEYS)
KEYS = list("BCDEFGHIJKLMNOPQRUV")


def _der(o: dict) -> dict:
    """Port of JS der(): derive D, L, Q, R from the base aggregates."""
    o["D"] = n(o.get("B")) + n(o.get("C"))
    o["L"] = o["D"] - (n(o.get("E")) + n(o.get("F")) + n(o.get("G")) + n(o.get("H"))
                       + n(o.get("I")) + n(o.get("J")) + n(o.get("K")))
    o["Q"] = o["D"] + n(o.get("N")) + n(o.get("O"))
    o["R"] = n(o.get("N")) + o["L"] + n(o.get("O")) - n(o.get("P"))
    return o


def _agg(emp_list: list[Employee], hold: set[str]) -> dict:
    """Port of JS agg()."""
    o = {k: 0 for k in "BCEFGHIJKMNOPUV"}
    o["K"] = 0
    o["P"] = 0
    o["M"] = 0
    for e in emp_list:
        o["U"] += e.epf
        o["V"] += e.eesic
        o["B"] += e.fixed
        o["C"] += e.arr
        o["E"] += e.tds
        o["F"] += e.pf
        o["G"] += e.esic
        o["H"] += e.loan + e.adv + e.oth
        o["J"] += e.nps
        o["M"] += 1
        o["N"] += e.reimb
        o["O"] += e.var
        if e.code in hold:
            o["K"] += e.net - e.var
            o["P"] += e.var
    o.setdefault("I", 0)
    return _der(o)


def _sumO(objs: list[dict], keys: list[str]) -> dict:
    o = {}
    for k in keys:
        o[k] = sum(n(x.get(k) if x else 0) for x in objs)
    return o


@dataclass
class Line:
    label: str
    kind: str                 # row | sum | memo | add
    o: dict
    id: int = -1
    ho: int = 0
    sales: int = 0
    t: int = 0                # total/bold row
    m: int = 0                # memo-block total
    frm: "Line | None" = None
    to: "Line | None" = None
    parts: list = field(default_factory=list)
    inp: list[str] = field(default_factory=list)


@dataclass
class Model:
    L: list[Line]
    warn: list[str]
    hoRows: list[Line]
    sRows: list[Line]
    totA: Line
    dir: Line
    totAB: Line
    grand: Line
    totS: Line
    jk: Line
    hra: Line
    hire: Line
    sv: Line
    e13: Line
    e325: Line
    totC: Line
    tot: Line
    dirs: list[Employee]
    ho: list[Employee]
    sl: list[Employee]


@dataclass
class ManualInputs:
    jk: float = 0.0
    jkm: float = 0.0
    hra: float = 0.0
    hire: float = 0.0
    dirc: float = 0.0


def build_model(ho_emps: list[Employee], sl_emps: list[Employee],
                hold: set[str], manual: ManualInputs) -> Model:
    """Port of JS model()."""
    dirs = [e for e in ho_emps if key(e.dept) == "director office"]
    rest = [e for e in ho_emps if e not in dirs]
    warn: list[str] = []

    names = list(HO_ORDER)
    for e in rest:
        if not any(key(x) == key(e.dept) for x in names):
            names.append(e.dept)
            warn.append(f'New HO department "{e.dept}" added as an extra row')

    L: list[Line] = []
    _counter = {"id": 0}

    def add(l: Line) -> Line:
        l.id = _counter["id"]
        _counter["id"] += 1
        L.append(l)
        return l

    hoRows = [add(Line(label=nm, kind="row",
                       o=_agg([e for e in rest if key(e.dept) == key(nm)], hold), ho=1))
              for nm in names]
    totA = add(Line(label="Total A", kind="sum", frm=hoRows[0], to=hoRows[-1],
                    o=_sumO([l.o for l in hoRows], KEYS), t=1))

    dO = _agg(dirs, hold)
    dO["C"] += n(manual.dirc)
    _der(dO)
    dir_line = add(Line(label="Directors (B)", kind="row", o=dO))

    totAB = add(Line(label="Total (A+B)", kind="sum", frm=totA, to=dir_line,
                     o=_sumO([totA.o, dir_line.o], KEYS), t=1))
    grand = totAB

    # Sales channels
    sn_defs = [
        ("Alternate Channels", lambda e: re.search(r"alternate", e.dept, re.I) is not None),
        ("GT", lambda e: re.search(r"general\s*trade|^gt$", e.dept, re.I) is not None),
    ]
    other = [e for e in sl_emps if not any(f(e) for _, f in sn_defs)]
    for d in _unique_in_order([e.dept for e in other]):
        sn_defs.append((d, (lambda dd: (lambda x: x.dept == dd))(d)))
        warn.append(f'New Sales department "{d}" added as an extra row')

    sRows = [add(Line(label=lb, kind="row",
                      o=_agg([e for e in sl_emps if f(e)], hold), sales=1))
             for lb, f in sn_defs]
    totS = add(Line(label="TOTAL (A)", kind="sum", frm=sRows[0], to=sRows[-1],
                    o=_sumO([l.o for l in sRows], KEYS), t=1))

    def memo(label: str, B: float, man: float) -> Line:
        o = {"B": B, "M": man}
        _der(o)
        return add(Line(label=label, kind="memo", o=o))

    jk, jkm, hra, hire = n(manual.jk), n(manual.jkm), n(manual.hra), n(manual.hire)
    J = memo("Promotor (JK)", jk, jkm)
    H = memo("HRA", hra, 0)
    Hi = memo("Hiring Fees @1000 per parameter", hire, 0)
    sv = memo("Service Charges @3%", round((jk + hra) * 0.03), 0)
    e13 = memo("EPF@13%", jk * 0.13, 0)
    e325 = memo("ESIC @3.25%", round(jk * 0.0325), 0)
    J.inp = ["B", "M"]
    H.inp = ["B"]
    Hi.inp = ["B"]

    totC = add(Line(label="TOTAL (B)", kind="sum", frm=J, to=e325,
                    o=_sumO([J.o, H.o, Hi.o, sv.o, e13.o, e325.o], KEYS), t=1, m=1))
    tot = add(Line(label="Total", kind="add", parts=[totS, totC],
                   o=_sumO([totS.o, totC.o], KEYS), t=1, m=1))

    return Model(L=L, warn=warn, hoRows=hoRows, sRows=sRows, totA=totA, dir=dir_line,
                 totAB=totAB, grand=grand, totS=totS, jk=J, hra=H, hire=Hi, sv=sv,
                 e13=e13, e325=e325, totC=totC, tot=tot, dirs=dirs, ho=ho_emps, sl=sl_emps)


def _unique_in_order(seq):
    seen = set()
    out = []
    for x in seq:
        if x not in seen:
            seen.add(x)
            out.append(x)
    return out


# ---------------------------------------------------------------------------
# Presentation column values (mirror JS leafVals / lineVals)
# ---------------------------------------------------------------------------

NC = list("BCDEFGHIJKLMNOPQRS")      # columns in the MIS report
MSET = list("BCEFGHIJKLMN")          # columns a memo / memo-total shows

NH = {"B": "Gross Earning", "C": "Extra/ Arrear", "D": "Variable", "E": "Gross Earning",
      "F": "TDS", "G": "EE PF", "H": "EE ESIC", "I": "Adv/Loan/Other Deduction",
      "J": "Medical", "K": "NPS", "L": "HOLD", "M": "Net Payout",
      "N": "No. Of Manpowers", "O": "Reimbursement Part", "P": "Total Gross",
      "Q": "Total Payout", "R": "ER PF", "S": "ER ESIC"}


def _leaf_vals(l: Line, mode: str) -> dict:
    o = l.o
    def gg(k):
        return n(o.get(k))
    B = gg("B"); C = gg("C"); D = gg("O"); F = gg("E"); G = gg("F"); H = gg("G")
    I = gg("H"); J = gg("I"); K = gg("J"); Lh = gg("K") + gg("P"); O = gg("N")
    E = (B + C) if (mode == "tpl" and l.sales) else (B + C + D)
    Mv = E - (F + G + H + I + J + K + Lh)
    P = (E + O + D) if mode == "tpl" else (E + O)
    return {"B": B, "C": C, "D": D, "E": E, "F": F, "G": G, "H": H, "I": I,
            "J": J, "K": K, "L": Lh, "M": Mv, "N": gg("M"), "O": O, "P": P,
            "Q": Mv + O, "R": gg("U"), "S": gg("V")}


def line_vals(M: Model, l: Line, mode: str, cache: dict) -> dict:
    if l.id in cache:
        return cache[l.id]
    if l.kind in ("row", "memo"):
        v = _leaf_vals(l, mode)
    else:
        if l.kind == "sum":
            kids = [x for x in M.L if l.frm.id <= x.id <= l.to.id]
        else:  # add
            kids = l.parts
        v = {}
        for k in NC:
            v[k] = sum(n(line_vals(M, x, mode, cache).get(k)) for x in kids)
    cache[l.id] = v
    return v


def cols_of(l: Line) -> list[str]:
    return MSET if (l.m or l.kind == "memo") else NC


def tpl_diff(M: Model) -> dict:
    a = line_vals(M, M.totAB, "tpl", {})
    b = line_vals(M, M.totAB, "fix", {})
    c = line_vals(M, M.totS, "tpl", {})
    d = line_vals(M, M.totS, "fix", {})
    return {"hoP": a["P"] - b["P"], "sQ": d["Q"] - c["Q"]}


# ---------------------------------------------------------------------------
# ESIC rule inference (mirror JS rdE / esicRule)
# ---------------------------------------------------------------------------

import math


def _rdE(v: float, mode: str) -> float:
    if float(v).is_integer():
        if mode == "ROUNDUP":
            return math.floor((v + 9999) / 10000)
        return math.floor((v + 5000) / 10000)
    if mode == "ROUNDUP":
        return math.ceil(v / 10000 - 1e-9)
    return math.floor(v / 10000 + 0.5 + 1e-9)


def esic_rule(E: list[Employee], bp: float, fld: str) -> dict:
    C = []
    for el in ("flag", "AC<21000"):
        for bs in ("AW", "AP+AT"):
            for m in ("ROUND", "ROUNDUP"):
                C.append({"el": el, "bs": bs, "m": m, "ok": 0})

    def fit(c, e):
        el = e.flag if c["el"] == "flag" else (e.ac < 21000)
        b = e.gross if c["bs"] == "AW" else (e.ap + e.at)
        return (_rdE(b * bp, c["m"]) if el else 0) == getattr(e, fld)

    for e in E:
        for c in C:
            if fit(c, e):
                c["ok"] += 1
    mx = max((c["ok"] for c in C), default=0)
    top = [c for c in C if c["ok"] == mx]
    mods = _unique_in_order([c["m"] for c in top])
    bad = [e for e in E if not any(fit(c, e) for c in top)]
    if len(top) == 1:
        t = top[0]
        desc = (f"{'ESIC flag = YES' if t['el']=='flag' else 'gross earnings < 21,000'}, "
                f"on {'gross earned' if t['bs']=='AW' else 'basic + special allowance'}, {t['m']}")
    else:
        desc = f"{len(top)} equivalent rules fit"
    return {"N": len(E), "top": top, "mods": mods, "bad": bad, "desc": desc}


# ---------------------------------------------------------------------------
# Cross-checks (mirror JS checks())
# ---------------------------------------------------------------------------

@dataclass
class Check:
    ok: bool
    t: str
    d: str = ""
    w: int = 0


def run_checks(M: Model, ho: ParsedSheet, sales: ParsedSheet, hold: set[str],
               manual: ManualInputs, mode: str, month_val: str) -> list[Check]:
    out: list[Check] = []
    all_e = list(M.ho) + list(M.sl)

    def ck(ok, t, d=""):
        out.append(Check(bool(ok), t, d))

    g1 = [e for e in all_e if abs(e.gross - (e.fixed + e.var + e.arr)) > 1]
    g2 = [e for e in all_e if abs(e.net - (e.gross - e.ded)) > 1]
    g3 = [e for e in all_e if abs(e.ded - (e.tds + e.pf + e.esic + e.loan + e.adv + e.oth + e.nps)) > 1]

    ck(not g1, "Gross = fixed earnings + variable + arrears (every employee)",
       (f"{len(g1)} mismatch: " + ", ".join(e.code for e in g1[:5])) if g1 else f"{len(all_e)} employees OK")
    ck(not g3, "Total deductions = sum of deduction columns",
       (f"{len(g3)} mismatch: " + ", ".join(e.code for e in g3[:5])) if g3 else "OK")
    ck(not g2, "Salary Net Pay = Gross Earning − Total Deduction (every employee)",
       (f"{len(g2)} mismatch: " + ", ".join(e.code for e in g2[:5])) if g2 else "OK")

    for parsed, lb in ((ho, "HO"), (sales, "Sales")):
        T = parsed.tot
        E = parsed.emps
        if not T:
            ck(False, lb + ": source total row not found, tie-out skipped", "")
            continue
        g = sum(e.gross for e in E)
        d = sum(e.ded for e in E)
        nt = sum(e.net for e in E)
        ok = abs(g - T["gross"]) <= 1 and abs(d - T["ded"]) <= 1 and abs(nt - T["net"]) <= 1
        ck(ok, f"{lb} sheet: employee rows tie to the sheet's own total row {T['row']} (no employee missed)",
           (f"gross {fmt_in(g)}, net pay {fmt_in(nt)}" if ok else
            f"rows: gross {fmt_in(g)}, net {fmt_in(nt)} vs total row: gross {fmt_in(T['gross'])}, net {fmt_in(T['net'])}"))
        eP = sum(e.epf for e in E)
        eE = sum(e.eesic for e in E)
        ok2 = ((T["epf"] is None or abs(eP - T["epf"]) <= 1)
               and (T["eesic"] is None or abs(eE - T["eesic"]) <= 1))
        if T["epf"] is not None or T["eesic"] is not None:
            ck(ok2, lb + " sheet: Employer PF / Employer ESIC totals tie to the sheet's total row",
               (f"{fmt_in(eP)} / {fmt_in(eE)}" if ok2 else
                f"employee rows {fmt_in(eP)} / {fmt_in(eE)} vs total row {fmt_in(T['epf'])} / {fmt_in(T['eesic'])} — a total-row formula may have been overwritten"))

    for parsed, lb in ((ho, "HO"), (sales, "Sales")):
        J = parsed.junk or []
        ck(not J, lb + ": no text typed into numeric columns",
           (f"{len(J)} cell(s): " + "; ".join(f"{j['code']} {NL.get(j['k'], j['k'])} = \"{j['v']}\"" for j in J[:6])
            + " — read as 0; restore the cell/formula in the salary sheet") if J else "OK")

    neg = [e for e in all_e if e.net < 0]
    ck(not neg, "No negative Net Pay", ", ".join(e.code for e in neg) if neg else "OK")

    seen = {}
    dup = []
    for e in all_e:
        if seen.get(e.code):
            dup.append(e)
        else:
            seen[e.code] = 1
    ck(not dup, "No duplicate employee codes", ", ".join(e.code for e in dup) if dup else "OK")

    G = _sumO([e.__dict__ for e in all_e], ["gross", "reimb", "net"])
    held = [e for e in all_e if e.code in hold]
    hs = sum(e.net for e in held)
    ck(abs(M.grand.o["Q"] + M.totS.o["Q"] - (G["gross"] + G["reimb"] + n(manual.dirc))) < 1,
       "MIS Total Gross = Σ source gross + reimbursement",
       f"{fmt_in(M.grand.o['Q'] + M.totS.o['Q'])} vs {fmt_in(G['gross'] + G['reimb'])}")
    ck(abs(M.grand.o["R"] + M.totS.o["R"] - (G["net"] + G["reimb"] - hs + n(manual.dirc))) < 1,
       "MIS Total Payout = Σ Net Pay + reimbursement − hold",
       f"{fmt_in(M.grand.o['R'] + M.totS.o['R'])} vs {fmt_in(G['net'] + G['reimb'] - hs)}")

    EE = _sumO([e.__dict__ for e in all_e], ["epf", "eesic"])
    ck(abs(M.totAB.o["U"] + M.totS.o["U"] - EE["epf"]) < 1 and abs(M.totAB.o["V"] + M.totS.o["V"] - EE["eesic"]) < 1,
       "ER PF / ER ESIC totals tie to the salary sheets (employer contributions)",
       f"{fmt_in(EE['epf'])} / {fmt_in(EE['eesic'])}")

    ck(True, f"Hold: {len(held)} employees, {fmt_in(hs)}", "see the Hold Review tab")

    for w in M.warn:
        ck(False, w, "")
    for parsed, lb in ((ho, "HO"), (sales, "Sales")):
        for w in parsed.warn:
            ck(False, (lb + ": ") + w, "")

    hi = [e for e in all_e if e.pf > 1800]
    if hi:
        out.append(Check(True, f"Info: PF above 1,800 for {len(hi)} employees — allowed by the "
                               "sheet's own PF formula (days 1–16 capped at 960 + days 17–30 "
                               "capped at 1,400); not an error", ""))

    for parsed, lb in ((ho, "HO"), (sales, "Sales")):
        E = parsed.emps
        if not E or not E[0].hasFlag:
            out.append(Check(True, lb + ": ESIC rule check skipped (ESIC flag / gross earnings column not found)", ""))
            continue
        for t, bp, f in (("Employee ESIC", 75, "esic"), ("Employer ESIC", 325, "eesic")):
            r = esic_rule(E, bp, f)
            ok = not r["bad"]
            rm = r["mods"][0] if len(r["mods"]) == 1 else "not distinguishable on this data"
            ck(ok, f"{lb}: {t} follows a consistent rule — {r['desc']}; rounding {rm}",
               (f"{r['N']} of {r['N']} employees match" if ok else
                f"{len(r['bad'])} employee(s) do not match ({', '.join(e.code for e in r['bad'][:5])}) — "
                "a manual override, or the sheet's ESIC rule has changed; send the updated sheet so the rule can be added"))
            if f == "eesic" and len(r["mods"]) == 1 and r["mods"][0] == "ROUNDUP" and (month_val or "") >= "2026-10":
                ck(False, lb + ": Employer ESIC values follow the old ROUNDUP rule, but the current "
                               "sheet rule is ROUND — the file may not have been recalculated before saving", "")

    if not n(manual.jk):
        out.append(Check(False, "Promotor (JK) not entered — Sales Total (B) and Total exclude it", "", w=1))

    if mode == "tpl":
        d = tpl_diff(M)
        out.append(Check(False,
                         f"Formula set = senior's template (as given): HO Total Gross is {fmt_in(d['hoP'])} "
                         "higher than payroll's gross + reimbursement (Variable counted twice); Sales Net "
                         f"Payout / Total Payout are {fmt_in(d['sQ'])} lower than payroll (Sales Gross Earning "
                         'leaves out Variable). Select "Corrected" to remove both.', "", w=1))
    return out


# ---------------------------------------------------------------------------
# Dashboard data (mirror JS dashData())
# ---------------------------------------------------------------------------

def dash_data(M: Model, ho: ParsedSheet, sales: ParsedSheet, hold: set[str], manual: ManualInputs) -> dict:
    G = M.totAB.o
    SS = M.totS.o

    def tot(k):
        return n(G.get(k)) + n(SS.get(k))

    def ln(l: Line):
        return {"label": l.label.replace(" (B)", ""), "M": n(l.o.get("M")), "B": n(l.o.get("B")),
                "Q": n(l.o.get("Q")), "R": n(l.o.get("R")), "K": n(l.o.get("K")) + n(l.o.get("P"))}

    ded = [["TDS", "E"], ["PF (employee)", "F"], ["ESIC (employee)", "G"],
           ["Loans / Advances / Other", "H"], ["NPS", "J"]]
    ded = [[a, tot(k)] for a, k in ded]

    held = [e for e in (list(M.ho) + list(M.sl)) if e.code in hold]
    by = {}
    for e in held:
        k = e.rem or "(none)"
        o = by.setdefault(k, [k, 0, 0])
        o[1] += 1
        o[2] += e.net

    ho_rows = [ln(l) for l in (M.hoRows + [M.dir]) if l.o.get("M")]
    return {"G": G, "SS": SS, "tot": tot, "ho": ho_rows, "sl": [ln(l) for l in M.sRows],
            "ded": ded, "gross": tot("D") + tot("O"),
            "hold": sorted(by.values(), key=lambda x: -x[2]), "heldN": len(held)}
