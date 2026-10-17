"""
Excel export for the APIS Salary MIS — ports exportX() and exportDash()
from the HTML tool, using openpyxl.

The MIS sheet is written with LIVE formulas (SUM ranges, per-row derivations,
memo-block percentage formulas) exactly as the JS builds them, so the delivered
workbook recalculates in Excel instead of carrying only static numbers.
"""

from __future__ import annotations

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import mis_core as C

NF = '#,##0;(#,##0);-'

_thin = Side(style="thin", color="FF808080")
_BD = Border(top=_thin, left=_thin, bottom=_thin, right=_thin)


def _col_idx(k: str) -> int:
    """MIS column letter (B..S) -> worksheet column number (B=2)."""
    return C.NC.index(k) + 2


def _fill(argb: str) -> PatternFill:
    return PatternFill(fill_type="solid", fgColor=argb)


def build_mis_workbook(M: C.Model, ho: C.ParsedSheet, sales: C.ParsedSheet,
                       hold: set[str], manual: C.ManualInputs, mode: str,
                       title: dict, checks: list[C.Check]) -> bytes:
    """Return .xlsx bytes for the full MIS workbook (MIS / HOLD&TDS / Checks)."""
    wb = Workbook()
    ws = wb.active
    ws.title = title["sheet"]

    cache: dict = {}

    def LVe(l):
        return C.line_vals(M, l, mode, cache)

    r = 1
    row_of: dict[int, int] = {}

    def hstyle(rr, c1, c2):
        for c in range(c1, c2 + 1):
            x = ws.cell(rr, c)
            x.font = Font(bold=True)
            x.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            x.border = _BD
            x.fill = _fill("FFD9E1F2")

    def header(lbl, first):
        nonlocal r
        ws.merge_cells(start_row=r, start_column=1, end_row=r + 1, end_column=1)
        ws.cell(r, 1).value = lbl
        ws.cell(r, 2).value = title["full"]
        hstyle(r, 1, 17 if first else 19)
        for k in C.NC:
            ws.cell(r + 1, _col_idx(k)).value = C.NH[k]
        hstyle(r + 1, 2, 19)
        for k in ("R", "S"):
            ws.cell(r + 1, _col_idx(k)).border = Border(left=_thin, right=_thin)
        ws.cell(r + 1, 1).border = Border(left=_thin, bottom=_thin, right=_thin)
        ws.row_dimensions[r + 1].height = 29
        r += 2

    header("HO & Directors", True)

    def put(l: C.Line):
        nonlocal r
        row_of[l.id] = r
        v = LVe(l)
        cs = C.cols_of(l)
        lb = bool(l.t) or l.label.startswith("Directors")
        if l is M.totA or l is M.totAB:
            bk = "BCDFGHIJKLMNOP"
        elif l is M.totS:
            bk = "BCDEFGHIJKLMNOP"
        elif l is M.totC or l is M.tot:
            bk = "BCEFGHIJKLMN"
        else:
            bk = ""

        def fx(k):
            if l.kind == "row":
                return {
                    "E": (f"B{r}+C{r}" if (mode == "tpl" and l.sales) else f"SUM(B{r}:D{r})"),
                    "M": f"E{r}-SUM(F{r}:L{r})",
                    "P": (f"E{r}+O{r}+D{r}" if mode == "tpl" else f"E{r}+O{r}"),
                    "Q": f"M{r}+O{r}",
                }.get(k)
            if l.kind == "memo":
                if k == "E":
                    return f"B{r}+C{r}"
                if k == "M":
                    return f"E{r}-SUM(F{r}:L{r})"
                if k == "B":
                    if l.label.startswith("Service"):
                        return f"ROUND((B{row_of[M.jk.id]}+B{row_of[M.hra.id]})*3%,0)"
                    if l.label == "EPF@13%":
                        return f"B{row_of[M.jk.id]}*13%"
                    if l.label.startswith("ESIC @"):
                        return f"ROUND(B{row_of[M.jk.id]}*3.25%,0)"
                return None
            if l.kind == "sum":
                a = row_of[l.frm.id]
                b = row_of[l.to.id]
                if not l.m and k == "E" and l is not M.totS:
                    return f"SUM(B{r}:D{r})"
                if not l.m and k == "Q":
                    return f"M{r}+O{r}"
                return f"SUM({k}{a}:{k}{b})"
            if l.kind == "add":
                return "+".join(k + str(row_of[p.id]) for p in l.parts)
            return None

        ws.cell(r, 1).value = l.label
        ws.cell(r, 1).border = _BD
        ws.cell(r, 1).font = Font(bold=lb)
        for k in cs:
            x = ws.cell(r, _col_idx(k))
            ff = fx(k)
            x.value = ("=" + ff) if ff else v[k]
            x.number_format = NF
            x.alignment = Alignment(horizontal="center")
            x.border = _BD
            x.font = Font(bold=(k in bk))
        for k in [("N" if kk == "M" else kk) for kk in (l.inp or [])]:
            if not v.get(k):
                ws.cell(r, _col_idx(k)).fill = _fill("FFFFFF00")
        r += 1

    for l in M.L:
        if l is M.sRows[0]:
            header("Sales", False)
        put(l)

    if not C.n(manual.jk):
        x = ws.cell(row_of[M.jk.id], 19)
        x.value = ("<< INPUT REQUIRED: Promotor (JK) / HRA / Hiring Fee not entered — "
                   "Totals (C) and Total exclude it." if mode == "tpl" else
                   "<< INPUT REQUIRED: Promotor (JK) / HRA / Hiring Fee not entered — "
                   "TOTAL (B) and Total exclude it.")
        x.font = Font(bold=True, color="FFFF0000")

    nt = ws.cell(r + 1, 1)
    nt.value = ("Rows EPF/ESIC Employer (Sales) are memo rows and, as in the approved MIS, "
                "are not included in TOTAL (C) / Total." if mode == "tpl" else
                "Total Gross = Gross Earning + Reimbursement; Total Payout = Net Payout + "
                "Reimbursement. ER PF / ER ESIC are employer contributions (not part of payout).")
    nt.font = Font(italic=True, size=9)

    ws.column_dimensions["A"].width = 34
    for k in C.NC:
        ws.column_dimensions[get_column_letter(_col_idx(k))].width = (
            16 if k in ("B", "E", "F") else 15)
    ws.freeze_panes = "B3"

    # ---- HOLD & TDS sheet ----
    hs = wb.create_sheet("HOLD&TDS")
    all_e = list(M.ho) + list(M.sl)
    q = {"r": 1}

    def block(ttl, rows_list, hd):
        hs.cell(q["r"], 1).value = ttl
        hs.cell(q["r"], 1).font = Font(bold=True, size=12)
        q["r"] += 1
        for i, h in enumerate(hd):
            x = hs.cell(q["r"], i + 1)
            x.value = h
            x.font = Font(bold=True)
            x.fill = _fill("FFD9E1F2")
        q["r"] += 1
        t = 0
        for i, a in enumerate(rows_list):
            vals = [i + 1] + list(a)
            for j, val in enumerate(vals):
                x = hs.cell(q["r"], j + 1)
                x.value = val
                if isinstance(val, (int, float)) and j > 3:
                    x.number_format = NF
            t += (a[3] if len(a) > 3 and isinstance(a[3], (int, float)) else 0)
            q["r"] += 1
        hs.cell(q["r"], 1).value = "Total"
        hs.cell(q["r"], 5).value = t
        hs.cell(q["r"], 5).number_format = NF
        hs.cell(q["r"], 1).font = Font(bold=True)
        hs.cell(q["r"], 5).font = Font(bold=True)
        q["r"] += 2

    block("Hold " + title["full"],
          [[e.code, e.name, e.dept, e.net, e.rem] for e in all_e if e.code in hold],
          ["S.No.", "Emp ID", "Name", "Dept.", "Net Pay Held", "Remarks"])
    block("TDS " + title["full"],
          [[e.code, e.name, e.dept, e.tds] for e in all_e if e.tds > 0],
          ["S.No.", "Emp ID", "Name", "Dept.", "Amount"])
    block("Deductions " + title["full"],
          [[e.code, e.name, e.dept, e.loan + e.adv + e.oth] for e in all_e if e.loan + e.adv + e.oth > 0],
          ["S.No.", "Emp ID", "Name", "Dept.", "Amount"])
    block("NPS " + title["full"],
          [[e.code, e.name, e.dept, e.nps] for e in all_e if e.nps > 0],
          ["S.No.", "Emp ID", "Name", "Dept.", "Amount"])
    for i, w in enumerate([8, 12, 30, 26, 16, 22]):
        hs.column_dimensions[get_column_letter(i + 1)].width = w

    # ---- Checks sheet ----
    cs = wb.create_sheet("Checks")
    cs.column_dimensions["A"].width = 100
    cs.cell(1, 1).value = "Cross-checks run at export — " + title["full"]
    cs.cell(1, 1).font = Font(bold=True, size=12)
    for i, c in enumerate(checks):
        cs.cell(i + 3, 1).value = ("OK  " if c.ok else "REVIEW  ") + c.t + (" — " + c.d if c.d else "")

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def build_dashboard_workbook(M: C.Model, ho: C.ParsedSheet, sales: C.ParsedSheet,
                             hold: set[str], manual: C.ManualInputs, title: dict) -> bytes:
    """Return .xlsx bytes for the Dashboard Summary workbook (ports exportDash)."""
    D = C.dash_data(M, ho, sales, hold, manual)
    G, SS, tot = D["G"], D["SS"], D["tot"]
    wb = Workbook()
    wb.remove(wb.active)

    def sh(name, head, rows, widths):
        ws = wb.create_sheet(name)
        hr = ws.append(head) or ws[ws.max_row]
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)
            c.fill = _fill("FFFBF0D4")
            c.border = Border(bottom=_thin)
        for row in rows:
            ws.append(row)
            for c in ws[ws.max_row]:
                if isinstance(c.value, (int, float)):
                    c.number_format = NF
        for i, w in enumerate(widths):
            ws.column_dimensions[get_column_letter(i + 1)].width = w

    jk_in = "Yes" if C.n(manual.jk) else "No — not entered"
    sh("Summary", ["Metric", title["full"]],
       [["Total Manpower", tot("M")], ["HO & Directors manpower", G["M"]],
        ["Sales manpower", SS["M"]], ["Total Gross", tot("Q")], ["Total Payout", tot("R")],
        ["HO & Directors Total Payout", G["R"]], ["Sales Total Payout", SS["R"]],
        *[[a, b] for a, b in D["ded"]],
        ["Hold (net pay withheld)", tot("K") + tot("P")], ["Employees on hold", D["heldN"]],
        ["Employer PF (ER PF)", tot("U")], ["Employer ESIC (ER ESIC)", tot("V")],
        ["Avg payout per head", round(tot("R") / tot("M")) if tot("M") else 0],
        ["Promotor (JK) included", jk_in]],
       [34, 20])

    sh("HO Departments",
       ["Department", "Manpower", "Net Earning", "Total Gross", "Total Payout", "Hold"],
       [*[[x["label"], x["M"], x["B"], x["Q"], x["R"], x["K"]] for x in D["ho"]],
        ["Total HO & Directors", G["M"], G["B"], G["Q"], G["R"], C.n(G.get("K")) + C.n(G.get("P"))]],
       [30, 12, 16, 16, 16, 14])

    sh("Sales",
       ["Channel", "Manpower", "Net Earning", "Total Gross", "Total Payout", "Hold"],
       [*[[x["label"], x["M"], x["B"], x["Q"], x["R"], x["K"]] for x in D["sl"]],
        ["Total (A)", SS["M"], SS["B"], SS["Q"], SS["R"], C.n(SS.get("K")) + C.n(SS.get("P"))]],
       [26, 12, 16, 16, 16, 14])

    sh("Deductions", ["Deduction", "Amount", "% of gross earning"],
       [[a, b, round(b / D["gross"] * 100, 2) if D["gross"] else 0] for a, b in D["ded"]],
       [30, 16, 18])

    sh("Hold by remark", ["Remarks", "Employees", "Net pay held"], D["hold"], [34, 12, 16])

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
